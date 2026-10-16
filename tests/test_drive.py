"""Drive flow against a fake Drive API (no network)."""
import io, re
import docx, httplib2, pytest
from googleapiclient.errors import HttpError
from fastapi.testclient import TestClient

from app import sources
from app.sources import parse_folder_id


def make_docx(lines, numbered):
    d = docx.Document()
    for l in lines:
        d.add_paragraph(l)
    for n in numbered:
        d.add_paragraph(n, style="List Number")
    b = io.BytesIO(); d.save(b); return b.getvalue()


ROOT = "ROOTFOLDER123456"
TREE = {
    ROOT: [
        {"id": "f314", "name": "ECEN314.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        {"id": "f214", "name": "ECEN214.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        {"id": "sub1", "name": "Archive", "mimeType": "application/vnd.google-apps.folder"},
        {"id": "sc1", "name": "shortcut to 350", "mimeType": "application/vnd.google-apps.shortcut",
         "shortcutDetails": {"targetId": "g350", "targetMimeType": "application/vnd.google-apps.document"}},
        {"id": "img", "name": "photo.png", "mimeType": "image/png"},
    ],
    "sub1": [{"id": "f303", "name": "ECEN303.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}],
}
CONTENT = {
    "f314": make_docx(["Learning Outcomes:", "Upon successfully completing the course, students will be able to:"],
                      ["Describe signals mathematically and understand how to perform mathematical operations on signals.",
                       "Identify commonly used signals such as the unit step.",
                       "Apply Fourier analysis to ideal low pass, high pass, band pass, and band-reject filtering."]),
    "f214": make_docx(["Learning Outcomes:", "Upon successfully completing the course, students will be able to:"],
                      ["Analyze the fundamental electrical quantities of charge, current, voltage, power, and energy."]),
    "f303": make_docx(["Learning Outcomes:"], ["Review basic notions of set theory.", "Explain sample spaces."]),
}
GDOCS = {"g350": "ECEN 350 Computer Architecture\n\nLearning Outcomes:\n1. Design a pipelined processor.\n2. Evaluate cache performance.\n"}
NAMES = {ROOT: "SimpleSyllabusAudit", "sub1": "Archive"}


def http_error(status):
    return HttpError(httplib2.Response({"status": status}), b"{}")


class Exec:
    def __init__(self, fn): self.fn = fn
    def execute(self): return self.fn()


class FakeFiles:
    def __init__(self, token): self.token = token
    def _ok(self):
        if self.token == "expired": raise http_error(401)
        if self.token != "good": raise http_error(404)
    def get(self, fileId, **kw):
        def run():
            self._ok()
            if fileId in NAMES: return {"id": fileId, "name": NAMES[fileId], "mimeType": "application/vnd.google-apps.folder"}
            if fileId == "f314": return {"id": fileId, "name": "ECEN314.docx", "mimeType": "x"}
            raise http_error(404)
        return Exec(run)
    def list(self, q, **kw):
        def run():
            self._ok()
            parents = re.findall(r"'([^']+)' in parents", q)
            items = [dict(f, modifiedTime="2026-09-01T00:00:00Z", webViewLink=f"https://drive/{f['id']}") for p in parents for f in TREE.get(p, [])]
            m = re.search(r'fullText contains \'"(.+?)"\'', q)
            if m:
                term = m.group(1).lower()
                items = [f for f in items if term in (GDOCS.get(f["id"], "") + str(CONTENT.get(f["id"], b""))).lower()]
            return {"files": items}
        return Exec(run)
    def export(self, fileId, mimeType):
        return Exec(lambda: (self._ok(), GDOCS[fileId].encode())[1])
    def get_media(self, fileId, **kw):
        return Exec(lambda: (self._ok(), CONTENT[fileId])[1])


class FakeSvc:
    def __init__(self, token): self._f = FakeFiles(token)
    def files(self): return self._f


@pytest.fixture
def client(monkeypatch, tmp_path):
    def fake_build(credentials=None, developerKey=None):
        return FakeSvc(getattr(credentials, "token", None) or developerKey)
    monkeypatch.setattr(sources, "_build", fake_build)
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    sources.Scanner._indexes.clear()
    from app import main
    monkeypatch.setattr(main, "text_cache", sources.TextCache(tmp_path / "c.json"))
    monkeypatch.setattr(main.llm, "API_KEY", None)
    return TestClient(main.app)


LINK = f"https://drive.google.com/drive/folders/{ROOT}?usp=sharing"


def test_parse_links():
    assert parse_folder_id(LINK) == ROOT
    assert parse_folder_id(f"https://drive.google.com/drive/u/1/folders/{ROOT}") == ROOT
    assert parse_folder_id(f"https://drive.google.com/open?id={ROOT}") == ROOT
    assert parse_folder_id(ROOT) == ROOT
    assert parse_folder_id("hello world") is None


def test_needs_signin_without_token(client):
    r = client.post("/api/folder", json={"folder_link": LINK}).json()
    assert r["ok"] is False and r["needs_signin"] is True


def test_no_access_and_expired(client):
    r = client.post("/api/folder", json={"folder_link": LINK}, headers={"X-Google-Token": "stranger"}).json()
    assert r["ok"] is False and not r.get("needs_signin")
    r = client.post("/api/folder", json={"folder_link": LINK}, headers={"X-Google-Token": "expired"}).json()
    assert r["ok"] is False and r["needs_signin"] is True


def test_connect_and_lookup(client):
    h = {"X-Google-Token": "good"}
    r = client.post("/api/folder", json={"folder_link": LINK}, headers=h).json()
    assert r["ok"] and r["folder_name"] == "SimpleSyllabusAudit" and r["files"] == 4 and r["access"] == "user"

    r = client.post("/api/lookup", json={"folder_link": LINK, "course_number": "ECEN 314"}, headers=h).json()
    assert r["status"] == "exists", r
    assert r["matches"][0]["name"] == "ECEN314.docx"
    items = r["objectives"]["items"]
    assert len(items) == 3 and items[0].startswith("Describe signals")

    r = client.post("/api/lookup", json={"folder_link": LINK, "course_number": "ecen303"}, headers=h).json()
    assert r["status"] == "exists" and r["matches"][0]["path"] == "Archive"

    r = client.post("/api/lookup", json={"folder_link": LINK, "course_number": "ECEN 350"}, headers=h).json()
    assert r["status"] == "exists" and r["objectives"]["items"] == ["Design a pipelined processor.", "Evaluate cache performance."]

    r = client.post("/api/lookup", json={"folder_link": LINK, "course_number": "ECEN 999"}, headers=h).json()
    assert r["status"] == "not_found" and r["files_scanned"] == 4


def test_other_user_cannot_reuse_cache(client):
    client.post("/api/folder", json={"folder_link": LINK}, headers={"X-Google-Token": "good"})
    r = client.post("/api/lookup", json={"folder_link": LINK, "course_number": "ECEN 314"},
                    headers={"X-Google-Token": "stranger"}).json()
    assert r["status"] == "error"


def test_empty_link_uses_samples(client):
    r = client.post("/api/folder", json={"folder_link": ""}).json()
    assert r["ok"] and r["access"] == "local" and r["files"] == 3


def test_folder_audit(client):
    import time
    h = {"X-Google-Token": "good"}
    r = client.post("/api/audit", json={"folder_link": LINK, "use_llm": False}, headers=h).json()
    assert r["ok"] and r["total"] == 4
    for _ in range(50):
        d = client.get(f"/api/audit/{r['job_id']}").json()
        if d["finished"]:
            break
        time.sleep(0.1)
    assert d["finished"] and d["summary"]["courses"] == 4
    by = {x["course"]: x for x in d["results"]}
    assert set(by) == {"ECEN 314", "ECEN 214", "ECEN 303", "ECEN 350"}
    assert by["ECEN 314"]["status"] == "avoid"            # "...and understand how to perform..."
    assert "understand" in by["ECEN 314"]["objectives"][0]["avoid_terms"]
    assert by["ECEN 350"]["status"] == "clean"
    assert by["ECEN 303"]["status"] == "clean"            # Review / Explain are both approved
    assert ("understand", 1) in [tuple(t) for t in d["summary"]["top_terms"]]


def test_audit_needs_signin(client):
    r = client.post("/api/audit", json={"folder_link": LINK}).json()
    assert r["ok"] is False and r["needs_signin"]


class _Creds:
    def __init__(self, token): self.token = token; self.service_account_email = "syllabus-app@example.iam.gserviceaccount.com"


def test_fixed_folder_no_signin(client, monkeypatch):
    monkeypatch.setenv("DRIVE_FOLDER_LINK", LINK)
    monkeypatch.setattr(sources, "_service_account_creds", lambda: _Creds("good"))
    cfg = client.get("/api/config").json()
    assert cfg["fixed_folder"] is True and cfg["google_client_id"] is None
    r = client.post("/api/folder", json={"folder_link": ""}).json()          # faculty send nothing
    assert r["ok"] and r["folder_name"] == "SimpleSyllabusAudit" and r["access"] == "service_account"
    r = client.post("/api/folder", json={"folder_link": "https://evil.example/other"}).json()   # client can't redirect it
    assert r["ok"] and r["folder_name"] == "SimpleSyllabusAudit"
    r = client.post("/api/lookup", json={"course_number": "ECEN 314"}).json()
    assert r["status"] == "exists"


def test_fixed_folder_not_shared(client, monkeypatch):
    monkeypatch.setenv("DRIVE_FOLDER_LINK", LINK)
    monkeypatch.setattr(sources, "_service_account_creds", lambda: _Creds("stranger"))
    r = client.post("/api/folder", json={"folder_link": ""}).json()
    assert r["ok"] is False and "syllabus-app@example.iam.gserviceaccount.com" in r["message"]
